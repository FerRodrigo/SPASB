from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from datetime import date
from .models import Entrada, Despesa
from .forms import EntradaForm, DespesaForm
from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from rest_framework.decorators import api_view
from rest_framework.response import Response
from .serializers import EntradaSerializer, DespesaSerializer


@login_required
def entrada_list(request):
    buscar = request.GET.get('buscar')
    tipo = request.GET.get('tipo')
    
    entradas = Entrada.objects.all().order_by('-data')
    
    if buscar:
        entradas = entradas.filter(descricao__icontains=buscar)
    if tipo:
        entradas = entradas.filter(tipo=tipo)
    
    total = entradas.aggregate(Sum('valor_total'))['valor_total__sum'] or 0
    total_ong = entradas.aggregate(Sum('valor_ong'))['valor_ong__sum'] or 0
    
    contexto = {'entradas': entradas, 'total': total, 'total_ong': total_ong, 'tipos_entrada': Entrada.TIPO_ENTRADA}
    return render(request, 'financeiro/entrada_list.html', contexto)


@login_required
def entrada_create(request):
    if request.method == 'POST':
        form = EntradaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Entrada cadastrada!')
            return redirect('financeiro:entrada_list')
    else:
        form = EntradaForm(initial={'data': date.today()})
    return render(request, 'financeiro/entrada_form.html', {'form': form, 'acao': 'Cadastrar'})


@login_required
def entrada_edit(request, pk):
    entrada = get_object_or_404(Entrada, pk=pk)
    if request.method == 'POST':
        form = EntradaForm(request.POST, instance=entrada)
        if form.is_valid():
            form.save()
            messages.success(request, 'Entrada atualizada!')
            return redirect('financeiro:entrada_list')
    else:
        form = EntradaForm(instance=entrada)
    return render(request, 'financeiro/entrada_form.html', {'form': form, 'acao': 'Editar', 'entrada': entrada})


@login_required
def entrada_delete(request, pk):
    entrada = get_object_or_404(Entrada, pk=pk)
    if request.method == 'POST':
        entrada.delete()
        messages.success(request, 'Entrada excluída!')
        return redirect('financeiro:entrada_list')
    return render(request, 'financeiro/entrada_confirm_delete.html', {'entrada': entrada})


@login_required
def despesa_list(request):
    buscar = request.GET.get('buscar')
    tipo = request.GET.get('tipo')
    
    despesas = Despesa.objects.all().order_by('-data')
    
    if buscar:
        despesas = despesas.filter(descricao__icontains=buscar)
    if tipo:
        despesas = despesas.filter(tipo=tipo)
    
    total = despesas.aggregate(Sum('valor'))['valor__sum'] or 0
    
    contexto = {'despesas': despesas, 'total': total, 'tipos_despesa': Despesa.TIPO_DESPESA}
    return render(request, 'financeiro/despesa_list.html', contexto)


@login_required
def despesa_create(request):
    if request.method == 'POST':
        form = DespesaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Despesa cadastrada!')
            return redirect('financeiro:despesa_list')
    else:
        form = DespesaForm(initial={'data': date.today()})
    return render(request, 'financeiro/despesa_form.html', {'form': form, 'acao': 'Cadastrar'})


@login_required
def despesa_edit(request, pk):
    despesa = get_object_or_404(Despesa, pk=pk)
    if request.method == 'POST':
        form = DespesaForm(request.POST, instance=despesa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Despesa atualizada!')
            return redirect('financeiro:despesa_list')
    else:
        form = DespesaForm(instance=despesa)
    return render(request, 'financeiro/despesa_form.html', {'form': form, 'acao': 'Editar', 'despesa': despesa})


@login_required
def despesa_delete(request, pk):
    despesa = get_object_or_404(Despesa, pk=pk)
    if request.method == 'POST':
        despesa.delete()
        messages.success(request, 'Despesa excluída!')
        return redirect('financeiro:despesa_list')
    return render(request, 'financeiro/despesa_confirm_delete.html', {'despesa': despesa})


@login_required
def relatorio_mensal(request):
    hoje = date.today()
    mes = int(request.GET.get('mes', hoje.month))
    ano = int(str(request.GET.get('ano', hoje.year)).replace('.', ''))
    
    entradas = Entrada.objects.filter(data__month=mes, data__year=ano)
    despesas = Despesa.objects.filter(data__month=mes, data__year=ano)
    
    total_entradas = entradas.aggregate(Sum('valor_total'))['valor_total__sum'] or 0
    total_ong = entradas.aggregate(Sum('valor_ong'))['valor_ong__sum'] or 0
    total_vet = entradas.aggregate(Sum('valor_veterinaria'))['valor_veterinaria__sum'] or 0
    total_despesas = despesas.aggregate(Sum('valor'))['valor__sum'] or 0
    saldo = total_ong - total_despesas
    
    contexto = {
        'entradas': entradas, 'despesas': despesas,
        'total_entradas': total_entradas, 'total_ong': total_ong,
        'total_vet': total_vet, 'total_despesas': total_despesas,
        'saldo': saldo, 'mes': mes, 'ano': ano,
    }
    return render(request, 'financeiro/relatorio_mensal.html', contexto)
@login_required
def relatorio_pdf(request):

    hoje = date.today()

    mes = int(request.GET.get('mes', hoje.month))
    ano = int(str(request.GET.get('ano', hoje.year)).replace('.', ''))

    entradas = Entrada.objects.filter(
        data__month=mes,
        data__year=ano
    ).order_by('data')

    despesas = Despesa.objects.filter(
        data__month=mes,
        data__year=ano
    ).order_by('data')

    total_entradas = entradas.aggregate(Sum('valor_total'))['valor_total__sum'] or 0
    total_ong = entradas.aggregate(Sum('valor_ong'))['valor_ong__sum'] or 0
    total_vet = entradas.aggregate(Sum('valor_veterinaria'))['valor_veterinaria__sum'] or 0
    total_despesas = despesas.aggregate(Sum('valor'))['valor__sum'] or 0
    saldo = total_ong - total_despesas

    meses = [
        'Janeiro', 'Fevereiro', 'Março', 'Abril',
        'Maio', 'Junho', 'Julho', 'Agosto',
        'Setembro', 'Outubro', 'Novembro', 'Dezembro'
    ]

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="Relatorio_SPASB_{mes:02d}_{ano}.pdf"'
    )

    doc = SimpleDocTemplate(
        response,
        pagesize=landscape(A4)
    )

    estilos = getSampleStyleSheet()

    titulo = ParagraphStyle(
        'Titulo',
        parent=estilos['Title'],
        alignment=TA_CENTER
    )

    elementos = []

    elementos.append(Paragraph("SPASB - Relatório Financeiro", titulo))
    elementos.append(Paragraph(f"{meses[mes-1]} de {ano}", estilos['Heading2']))
    elementos.append(Spacer(1, 20))

    resumo = [
        ['Indicador', 'Valor'],
        ['Entradas', f'R$ {total_entradas}'],
        ['Parte da ONG', f'R$ {total_ong}'],
        ['Veterinária', f'R$ {total_vet}'],
        ['Despesas', f'R$ {total_despesas}'],
        ['Saldo', f'R$ {saldo}'],
    ]

    tabela = Table(resumo)

    tabela.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.darkblue),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('BACKGROUND', (0,1), (-1,-1), colors.beige),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
    ]))

    elementos.append(tabela)
    elementos.append(Spacer(1, 20))

    elementos.append(Paragraph("Entradas", estilos['Heading2']))

    dados_entradas = [['Data', 'Descrição', 'Tipo', 'Total', 'ONG', 'Veterinária']]

    for e in entradas:
        dados_entradas.append([
            e.data.strftime('%d/%m/%Y'),
            escape(e.descricao),
            e.get_tipo_display(),
            f'R$ {e.valor_total}',
            f'R$ {e.valor_ong}',
            f'R$ {e.valor_veterinaria}',
        ])

    if len(dados_entradas) == 1:
        dados_entradas.append(['-', 'Nenhuma entrada', '-', '-', '-', '-'])

    tabela_e = Table(dados_entradas)
    tabela_e.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.green),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
    ]))

    elementos.append(tabela_e)
    elementos.append(Spacer(1, 20))

    elementos.append(Paragraph("Despesas", estilos['Heading2']))

    dados_despesas = [['Data', 'Descrição', 'Tipo', 'Valor']]

    for d in despesas:
        dados_despesas.append([
            d.data.strftime('%d/%m/%Y'),
            escape(d.descricao),
            d.get_tipo_display(),
            f'R$ {d.valor}',
        ])

    if len(dados_despesas) == 1:
        dados_despesas.append(['-', 'Nenhuma despesa', '-', '-'])

    tabela_d = Table(dados_despesas)
    tabela_d.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.red),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
    ]))

    elementos.append(tabela_d)

    doc.build(elementos)

    return response

@login_required
def relatorio_excel(request):

    hoje = date.today()

    mes = int(request.GET.get('mes', hoje.month))
    ano = int(str(request.GET.get('ano', hoje.year)).replace('.', ''))

    entradas = Entrada.objects.filter(
        data__month=mes,
        data__year=ano
    ).order_by('data')

    despesas = Despesa.objects.filter(
        data__month=mes,
        data__year=ano
    ).order_by('data')

    total_entradas = entradas.aggregate(
        Sum('valor_total')
    )['valor_total__sum'] or 0

    total_ong = entradas.aggregate(
        Sum('valor_ong')
    )['valor_ong__sum'] or 0

    total_vet = entradas.aggregate(
        Sum('valor_veterinaria')
    )['valor_veterinaria__sum'] or 0

    total_despesas = despesas.aggregate(
        Sum('valor')
    )['valor__sum'] or 0

    saldo = total_ong - total_despesas

    meses = [
        'Janeiro', 'Fevereiro', 'Março', 'Abril',
        'Maio', 'Junho', 'Julho', 'Agosto',
        'Setembro', 'Outubro', 'Novembro', 'Dezembro'
    ]

    # Criar arquivo Excel
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Relatório Mensal'

    # Título
    sheet['A1'] = 'SPASB - Relatório Financeiro'
    sheet['A1'].font = Font(
        bold=True,
        size=16
    )

    sheet['A2'] = f'{meses[mes - 1]} de {ano}'
    sheet['A2'].font = Font(
        bold=True,
        size=12
    )

    # Resumo
    sheet['A4'] = 'RESUMO FINANCEIRO'
    sheet['A4'].font = Font(
        bold=True,
        size=12
    )

    resumo = [
        ('Total de entradas', total_entradas),
        ('Parte da ONG', total_ong),
        ('Parte da veterinária', total_vet),
        ('Total de despesas', total_despesas),
        ('Saldo da ONG', saldo),
    ]

    linha = 5

    for descricao, valor in resumo:

        sheet.cell(
            row=linha,
            column=1,
            value=descricao
        )

        sheet.cell(
            row=linha,
            column=2,
            value=float(valor)
        )

        sheet.cell(
            row=linha,
            column=2
        ).number_format = 'R$ #,##0.00'

        linha += 1

    # Entradas
    linha += 1

    sheet.cell(
        row=linha,
        column=1,
        value='ENTRADAS'
    )

    sheet.cell(
        row=linha,
        column=1
    ).font = Font(
        bold=True,
        size=12
    )

    linha += 1

    cabecalho_entradas = [
        'Data',
        'Descrição',
        'Tipo',
        'Valor Total',
        'ONG',
        'Veterinária'
    ]

    for coluna, titulo in enumerate(cabecalho_entradas, start=1):

        celula = sheet.cell(
            row=linha,
            column=coluna,
            value=titulo
        )

        celula.font = Font(bold=True)
        celula.fill = PatternFill(
            'solid',
            fgColor='198754'
        )
        celula.font = Font(
            bold=True,
            color='FFFFFF'
        )
        celula.alignment = Alignment(
            horizontal='center'
        )

    linha += 1

    for entrada in entradas:

        sheet.cell(
            row=linha,
            column=1,
            value=entrada.data
        )

        sheet.cell(
            row=linha,
            column=1
        ).number_format = 'dd/mm/yyyy'

        sheet.cell(
            row=linha,
            column=2,
            value=entrada.descricao
        )

        sheet.cell(
            row=linha,
            column=3,
            value=entrada.get_tipo_display()
        )

        sheet.cell(
            row=linha,
            column=4,
            value=float(entrada.valor_total)
        )

        sheet.cell(
            row=linha,
            column=5,
            value=float(entrada.valor_ong)
        )

        sheet.cell(
            row=linha,
            column=6,
            value=float(entrada.valor_veterinaria)
        )

        for coluna in range(4, 7):
            sheet.cell(
                row=linha,
                column=coluna
            ).number_format = 'R$ #,##0.00'

        linha += 1

    # Despesas
    linha += 2

    sheet.cell(
        row=linha,
        column=1,
        value='DESPESAS'
    )

    sheet.cell(
        row=linha,
        column=1
    ).font = Font(
        bold=True,
        size=12
    )

    linha += 1

    cabecalho_despesas = [
        'Data',
        'Descrição',
        'Tipo',
        'Valor'
    ]

    for coluna, titulo in enumerate(cabecalho_despesas, start=1):

        celula = sheet.cell(
            row=linha,
            column=coluna,
            value=titulo
        )

        celula.font = Font(
            bold=True,
            color='FFFFFF'
        )

        celula.fill = PatternFill(
            'solid',
            fgColor='DC3545'
        )

        celula.alignment = Alignment(
            horizontal='center'
        )

    linha += 1

    for despesa in despesas:

        sheet.cell(
            row=linha,
            column=1,
            value=despesa.data
        )

        sheet.cell(
            row=linha,
            column=1
        ).number_format = 'dd/mm/yyyy'

        sheet.cell(
            row=linha,
            column=2,
            value=despesa.descricao
        )

        sheet.cell(
            row=linha,
            column=3,
            value=despesa.get_tipo_display()
        )

        sheet.cell(
            row=linha,
            column=4,
            value=float(despesa.valor)
        )

        sheet.cell(
            row=linha,
            column=4
        ).number_format = 'R$ #,##0.00'

        linha += 1

    # Ajustar largura das colunas
    larguras = {
        'A': 15,
        'B': 35,
        'C': 25,
        'D': 18,
        'E': 18,
        'F': 18,
    }

    for coluna, largura in larguras.items():
        sheet.column_dimensions[coluna].width = largura

    # Congelar cabeçalho inicial
    sheet.freeze_panes = 'A4'

    # Preparar resposta
    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument'
            '.spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        f'attachment; filename="Relatorio_SPASB_{mes:02d}_{ano}.xlsx"'
    )

    workbook.save(response)

    return response

@api_view(['GET'])
def api_entradas(request):
    entradas = Entrada.objects.all().order_by('-data')
    serializer = EntradaSerializer(entradas, many=True)
    return Response(serializer.data)

@api_view(['GET'])
def api_despesas(request):
    despesas = Despesa.objects.all().order_by('-data')
    serializer = DespesaSerializer(despesas, many=True)
    return Response(serializer.data)